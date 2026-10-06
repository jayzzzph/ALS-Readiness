import asyncio
import sys

if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())

from pathlib import Path

from app.core.security import generate_temp_password, hash_password
from app.db.session import AsyncSessionLocal
from app.models.learner import Learner
from app.models.learning_strand import LearningStrand
from app.models.lri_test import LRITest
from app.models.lri_test_attempt import LRITestAttempt
from app.models.participant_intake import ParticipantIntake
from app.models.strand_test import StrandTest
from app.models.strand_test_attempt import StrandTestAttempt
from app.models.user import User
from app.models.user_profile import UserProfile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

BASE_DIR = Path(__file__).resolve().parent.parent.parent
DATA_DIR = BASE_DIR / "data"

workbook = load_workbook(
    DATA_DIR / "pretest_data.xlsx",
    read_only=True,
    data_only=True,
)

participant_intake_sheet = workbook["Participant Intakes"]
lri_test_summary = workbook["LRI Test Summary"]
diagnostic_test_summary = workbook["Diagnostic Test Summary"]

consolidated_data = zip(
    participant_intake_sheet.iter_rows(min_row=2, values_only=True),
    lri_test_summary.iter_rows(min_row=2, values_only=True),
    diagnostic_test_summary.iter_rows(min_row=2, values_only=True)
)

participants = []

for pi, lri, dt  in consolidated_data:
    data = {
        "id_no": pi[0],
        "first_name": pi[1],
        "last_name": pi[2],
        "age": int(pi[3]),
        "sex": "male" if pi[4].lower() == "m" else "female",
        "civil_status": pi[5],
        "highest_educational_attainment": pi[6],
        "als_learning_strands": pi[7].split(", "),
        "als_enrollment_months": int(pi[8]),
        "has_taken_ae_test": bool(pi[9]),
        "lri_test": {
            "lri_score": float(lri[1])
        },
        "diagnostic_test": {
            "ls1_en_score": float(dt[1]),
            "ls1_fil_score": float(dt[2]),
            "ls3_score": float(dt[3]),
        },
    }

    participants.append(data)


async def seed_participants(session: AsyncSession) -> None:
    lri_test = (
        await session.execute(
            select(LRITest)
        )
    ).scalar_one_or_none()

    diagnostic_tests = (
        await session.execute(
            select(StrandTest, LearningStrand.code)
            .join(LearningStrand, StrandTest.strand_id == LearningStrand.id)
            .where(StrandTest.type == "pretest")
        )
    ).all()

    ls1_en_test = next(
        (test[0] for test in diagnostic_tests if test[1] == "LS1-EN"), 
        None
    )

    ls1_fil_test = next(
        (test[0] for test in diagnostic_tests if test[1] == "LS1-FIL"), 
        None
    )

    ls3_test = next(
        (test[0] for test in diagnostic_tests if test[1] == "LS3"), 
        None
    )

    for participant in participants:
        # Creat user entity
        temp_pass = generate_temp_password()

        user = User(
            id_no=participant.get("id_no"),
            password_hash=hash_password(temp_pass),
        )

        session.add(user)
        await session.flush()
        await session.refresh(user)

        # Create learner entity
        learner = Learner(user_id=user.id)

        session.add(learner)
        await session.flush()
        await session.refresh(learner)

        # Create user profile
        user_profile = UserProfile(
            user_id=user.id,
            first_name=participant.get("first_name"),
            last_name=participant.get("last_name"),
        )

        session.add(user_profile)

        # Participant Intake
        participant_intake = ParticipantIntake(
            user_id=user.id,
            age=participant.get("age"),
            sex=participant.get("sex"),
            civil_status=participant.get("civil_status"),
            highest_educational_attainment=participant.get("highest_educational_attainment"),
            als_learning_strands=participant.get("als_learning_strands"),
            als_enrollment_months=participant.get("als_enrollment_months"),
            has_taken_ae_test=participant.get("has_taken_ae_test"),
            ae_test_attempt_count=0,
        )

        session.add(participant_intake)

        # Create lri test attempt
        lri_data = participant.get("lri_test")
        lri_test_attempt = LRITestAttempt(
            test_id=lri_test.id,
            learner_id=learner.id,
            lri_score=lri_data.get("lri_score"),
        )

        session.add(lri_test_attempt)

        # Crete diagnostic test
        diagnostic_data = participant.get("diagnostic_test")

        ls1_en_attempt = StrandTestAttempt(
            test_id=ls1_en_test.id,
            learner_id=learner.id,
            total_score=diagnostic_data.get("ls1_en_score"),
            item_count=20,
        )

        ls1_fil_attempt = StrandTestAttempt(
            test_id=ls1_fil_test.id,
            learner_id=learner.id,
            total_score=diagnostic_data.get("ls1_fil_score"),
            item_count=20,
        )

        ls3_attempt = StrandTestAttempt(
            test_id=ls3_test.id,
            learner_id=learner.id,
            total_score=diagnostic_data.get("ls3_score"),
            item_count=20,
        )

        session.add(ls1_en_attempt)
        session.add(ls1_fil_attempt)
        session.add(ls3_attempt)\
        
        await session.flush()
        await session.refresh(participant_intake)
        await session.refresh(lri_test_attempt)
        await session.refresh(ls1_en_attempt)
        await session.refresh(ls1_fil_attempt)
        await session.refresh(ls3_attempt)
        await session.refresh(user_profile)

        print(f"{user_profile.first_name} {user_profile.last_name} - id_no: {user.id_no}, password: {temp_pass}")


async def main() -> None:
    async with AsyncSessionLocal() as session, session.begin():
        await seed_participants(session)


if __name__ == "__main__":
    asyncio.run(main())
